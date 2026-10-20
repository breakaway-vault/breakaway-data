#!/usr/bin/env python3
"""
breakaway-data: a preservation archive of NHL Breakaway highlight metadata.

USAGE
  python breakaway.py all                 # fetch (only new/missing) -> covers -> build -> package
  python breakaway.py fetch [--refresh] [--ids 12,40] [--discover]
                                          # highlight->token links + raw nft. metadata for the card list
  python breakaway.py covers              # full-res cover art into output/covers_full/ (sha256-checked)
  python breakaway.py build               # OFFLINE: rebuild output/highlights.xlsx from data/ only
  python breakaway.py package             # output/covers/, covers.zip, Breakaway-Highlights.zip
  python breakaway.py videos [--out-dir DIR] [--limit N] [--series "Series 3"] [--ids 12,40] [--yes]
                                          # OPT-IN video download (asks y/n after a size estimate)
  python breakaway.py verify              # re-hash data/meta, covers and any videos against stored sha256
  python breakaway.py add <url|seriesId/tokenId> ...   # add moments (also reads data/extra_moments.txt)
  python breakaway.py album [--layout 1up|2x2|3x3] [--page letter|a4] [--set ..] [--series ..] [--player ..]
                            [--team ..] [--bleed] [--out file.pdf] [--max-px N --jpeg Q]   # printable PDF album
  python breakaway.py view                # offline viewer/index.html (search, filters, details, video)
  IPFS: --gateway URL (repeatable) adds gateways tried first; --ipfs-only skips the Breakaway CDN.
  Common flags: --concurrency N (default 6, max 8) for nft.nhlbreakaway.com / CDN requests.

SOURCES (all public, no login)
  card list     data/highlights.json (committed), plus moments added with `add` / data/extra_moments.txt,
                plus (with `fetch --discover`) moments NHL Breakaway's own set listings show that aren't in it yet
  links         NHL Breakaway set listings (apigw.sweet.io/catalog/search/moment/set/<set>), matched on
                set + player + series; existing links in data/links.json are kept as they are
  on-chain-ish  https://nft.nhlbreakaway.com/series/<series>/<token>  saved byte-for-byte to data/meta/<id>.json
  off-chain     raw set listings + set completion + quest data in data/offchain/
  covers/video  nft.nhlbreakaway.com image CDN and mediagateway.sweet.io (output/, gitignored)

IDS
  "id" in data/highlights.json is the card id from the original card list (1..1818, not contiguous). Moments
  added later (`add`, `fetch --discover`) get the next free numbers and "on_site": false.
"""
import argparse, base64, concurrent.futures as cf, csv, datetime as dt, hashlib, html, io, json, os, re
import sys, threading, time, unicodedata, urllib.error, urllib.request, zipfile

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA, OUT, CACHE = (os.path.join(ROOT, d) for d in ("data", "output", "cache"))
NFT = "https://nft.nhlbreakaway.com/series"
APIGW = "https://apigw.sweet.io"
UA = "Mozilla/5.0 (breakaway-data preservation export; polite, cached)"
HDR = {"User-Agent": UA, "Referer": "https://nhlbreakaway.com/", "Origin": "https://nhlbreakaway.com"}


def p(*a): print(*a, flush=True)
def idkey(x):
    x = str(x[0] if isinstance(x, tuple) else x)
    return (0, int(x), "") if x.isdigit() else (1, 0, x)
def now(): return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
def jload(path, default=None):
    return json.load(open(path, encoding="utf-8")) if os.path.exists(path) else default
def jdump(obj, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"; json.dump(obj, open(tmp, "w", encoding="utf-8"), indent=1, ensure_ascii=False); os.replace(tmp, path)
def wbytes(path, b):
    os.makedirs(os.path.dirname(path), exist_ok=True); tmp = path + ".part"
    open(tmp, "wb").write(b); os.replace(tmp, path)


# ------------------------------------------------------------------ polite HTTP
class Throttle:
    def __init__(self, n, base=0.15):
        self.sem = threading.Semaphore(n); self.base = base; self.delay = base; self.lock = threading.Lock()
    def ok(self):
        with self.lock: self.delay = max(self.base, self.delay * 0.9)
    def bad(self):
        with self.lock: self.delay = min(30.0, self.delay * 2 + 0.5)

THR = Throttle(6)


def http(url, method="GET", tries=6, timeout=60, thr=None, headers=None):
    """Returns (status, bytes, headers). 401/403/404 are final (never worked around)."""
    thr = thr or THR
    for i in range(tries):
        thr.sem.acquire(); time.sleep(thr.delay)
        try:
            req = urllib.request.Request(url, headers={**HDR, **(headers or {})}, method=method)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                b = r.read() if method == "GET" else b""
                thr.ok(); return r.status, b, r.headers
        except urllib.error.HTTPError as e:
            if e.code in (400, 401, 403, 404, 410):
                return e.code, None, e.headers
            thr.bad(); ra = e.headers.get("Retry-After") if e.headers else None
            time.sleep(float(ra) if ra and ra.isdigit() else min(60, 2 ** (i + 1)))
        except Exception:
            thr.bad(); time.sleep(min(60, 2 ** (i + 1)))
        finally:
            thr.sem.release()
    return None, None, None


def norm(s):
    s = unicodedata.normalize("NFKD", html.unescape(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", s)
def strip(s): return html.unescape(re.sub(r"<[^>]+>", "", s or "")).strip()


# ------------------------------------------------------------------ card list
def load_cards():
    """The card list is data/highlights.json (it ships with the repo). New moments come from `add`,
    data/extra_moments.txt or `fetch --discover`; nothing is scraped from any third-party site."""
    rows = jload(os.path.join(DATA, "highlights.json"))
    if not rows: sys.exit("data/highlights.json is missing or empty. Restore it from the repo, then re-run.")
    rows.sort(key=lambda r: idkey(r["id"]))
    return rows


# ------------------------------------------------------------------ off-chain set data + links
def fetch_offchain(refresh):
    od = os.path.join(DATA, "offchain")
    thr = Throttle(3, 0.3)
    def grab(rel, url):
        path = os.path.join(od, rel)
        if refresh or not os.path.exists(path):
            st, b, _ = http(url, thr=thr)
            if b: wbytes(path, b)
        return jload(path, [])
    sc = grab("setcompletion_826.json", f"{APIGW}/gameplay/setcompletion?campaign_id=826")
    cc = grab("collection_challenge_826.json", f"{APIGW}/gameplay/collection-challenge?campaign_id=826")
    sets = {x["set_uuid"] for x in sc if x.get("set_uuid")}
    for c in cc:
        for h in c.get("highlights") or []:
            if h.get("moment_uuid"): sets.add(h["moment_uuid"][:12] + "0" * 20)
    md = os.path.join(DATA, "meta")
    if os.path.isdir(md):
        for fn in os.listdir(md):
            try:
                a = {t["name"]: t["value"] for t in jload(os.path.join(md, fn)).get("attributes", [])}
                if a.get("Set UUID"): sets.add(a["Set UUID"])
            except Exception: pass
    with cf.ThreadPoolExecutor(3) as ex:
        list(ex.map(lambda u: grab(f"sets/{u}.json", f"{APIGW}/catalog/search/moment/set/{u}?campaign_id=826"), sorted(sets)))
    moments = []
    for fn in sorted(os.listdir(os.path.join(od, "sets"))):
        for m in jload(os.path.join(od, "sets", fn), []):
            sc_ = re.search(r"/series/([A-Za-z0-9]{8})/", json.dumps(m.get("assets")))
            moments.append({"uuid": m["uuid"], "player": m.get("name"), "set": (m.get("token_set") or {}).get("name"),
                            "series_name": (m.get("token_series") or {}).get("name"),
                            "rarity": (m.get("rarity") or {}).get("name"), "series_id": sc_.group(1) if sc_ else None,
                            "file": f"data/offchain/sets/{fn}", "raw": m})
    p(f"Off-chain: {len(sets)} sets, {len(moments)} moments listed by NHL Breakaway")
    return moments


SUF = re.compile(r"\s*\((S1|S2|S3|I|II|III)\)\s*$")
SUF_SERIES = {"S1": "Series 1", "I": "Series 1", "S2": "Series 2", "II": "Series 2", "S3": "Series 3", "III": "Series 3"}


def build_links(rows, moments):
    path = os.path.join(DATA, "links.json"); links = jload(path, {})
    idx = {}
    for m in moments: idx.setdefault((norm(m["set"]), norm(m["player"])), {})[m["uuid"]] = m
    todo = []
    for r in rows:
        k = str(r["id"])
        if k in links: continue
        sm = SUF.search(r["set"]); want = SUF_SERIES[sm.group(1)] if sm else None
        cands = list(idx.get((norm(SUF.sub("", r["set"])), norm(r["player"])), {}).values())
        if want: cands = [m for m in cands if m["series_name"] == want]
        elif len({m["series_name"] for m in cands}) > 1: cands = [m for m in cands if m["series_name"] == "Series 1"] or cands
        if len(cands) > 1: cands = [m for m in cands if norm(m["rarity"]) == norm(r["rarity"])] or cands
        if len(cands) == 1 and cands[0]["series_id"]:
            m = cands[0]
            links[k] = {"series_id": m["series_id"], "token_id": str(int(m["uuid"][16:20], 16) << 48),
                        "moment_uuid": m["uuid"], "via": "NHL Breakaway set listing match (set+player+series)"}
        else: todo.append(r["id"])
    if todo:
        p(f"{len(todo)} highlight(s) could not be matched to an NHL Breakaway set listing: "
          f"{', '.join(map(str, todo[:10]))}{' ...' if len(todo) > 10 else ''}. "
          "Add them by URL with `python breakaway.py add <nhlbreakaway.com URL>`.")
    jdump(dict(sorted(links.items(), key=idkey)), path)
    return links


# ------------------------------------------------------------------ raw nft. metadata
def b64sha(b): return base64.urlsafe_b64encode(hashlib.sha256(b).digest()).decode()


def ipfs_entries(d):
    """Every (path, ipfs CID, sha256, url) found anywhere in a metadata doc."""
    out = []
    def walk(o, path):
        if isinstance(o, dict):
            if isinstance(o.get("altloc"), dict) and o["altloc"].get("ipfs"):
                out.append({"path": path, "ipfs": o["altloc"]["ipfs"], "sha256": o.get("sha256"), "url": o.get("url") or o.get("uri")})
            for k, v in o.items(): walk(v, f"{path}.{k}" if path else k)
        elif isinstance(o, list):
            for i, v in enumerate(o): walk(v, f"{path}[{i}]")
    walk(d, ""); return out


def get_nft(sid, tok, exact=False):
    """GET nft.nhlbreakaway.com metadata. Edition #1 of a moment can be missing (never minted / burned);
    unless the token is exact, try the next few serials of the same moment (one request each)."""
    tok = int(tok); tries = [tok] if exact or tok & 0xFFFFFFFFFFFF else [tok + n for n in range(4)]
    st = b = h = None
    for t in tries:
        st, b, h = http(f"{NFT}/{sid}/{t}")
        if b or st != 404: return st, b, h, str(t)
    return st, b, h, str(tries[-1])


def fetch_meta(links, refresh, only=None):
    mdir = os.path.join(DATA, "meta"); mpath = os.path.join(DATA, "manifest.json")
    man = jload(mpath, {}); lock = threading.Lock()
    def one(k):
        ln = links[k]; url = f"{NFT}/{ln['series_id']}/{ln['token_id']}"; fp = os.path.join(mdir, f"{k}.json")
        e = man.get(k, {})
        if not refresh and os.path.exists(fp) and e.get("status") == 200 and e.get("source_url") == url:
            return
        st, b, h, tok = get_nft(ln["series_id"], ln["token_id"], ln.get("exact"))
        if b and tok != str(ln["token_id"]):  # edition #1 missing: keep the exact token that answered
            with lock: links[k] = ln = {**ln, "token_id": tok, "exact": True,
                                        "via": f"{ln.get('via', '')} (edition #1 token missing; serial {int(tok) & 0xFFFF} used)".strip()}
            url = f"{NFT}/{ln['series_id']}/{tok}"
        e = {"id": int(k) if k.isdigit() else k, "source_url": url, "fetched_at": now(), "status": st, "link_via": ln.get("via")}
        if b:
            wbytes(fp, b)
            e.update({"meta_path": f"data/meta/{k}.json", "bytes": len(b), "sha256": hashlib.sha256(b).hexdigest(),
                      "content_type": h.get("Content-Type") if h else None})
            try: e["ipfs"] = ipfs_entries(json.loads(b))
            except Exception: e["parse_error"] = True
        with lock: man[k] = e
    keys = sorted((k for k in links if only is None or k in only), key=idkey)
    with cf.ThreadPoolExecutor(THR.sem._value) as ex:
        for n, _ in enumerate(ex.map(one, keys), 1):
            if n % 250 == 0: p(f"  metadata {n}/{len(keys)}")
    for _ in range(2):  # retry failures, slower
        bad = [k for k in keys if man.get(k, {}).get("status") != 200]
        if not bad: break
        p(f"  retrying {len(bad)} failed metadata fetches"); time.sleep(10)
        for k in bad: one(k)
    jdump(dict(sorted(man.items(), key=idkey)), mpath)
    jdump(dict(sorted(links.items(), key=idkey)), os.path.join(DATA, "links.json"))
    return man


def discover_refs(rows, moments, man):
    """Moments in NHL Breakaway's public set listings that aren't in the card list yet -> (series, token)."""
    known = {e.get("moment_uuid") for e in man.values()} | {v.get("moment_uuid") for v in jload(os.path.join(DATA, "links.json"), {}).values()}
    out = []
    for m in moments:
        if m["uuid"] in known or not m["series_id"]: continue
        out.append((m["series_id"], str(int(m["uuid"][16:20], 16) << 48))); known.add(m["uuid"])
    return out


def cmd_fetch(a):
    rows = load_cards()
    only = {x.strip() for x in a.ids.split(",") if x.strip()} if a.ids else None
    if only:
        unknown = only - {str(r["id"]) for r in rows}
        if unknown: sys.exit(f"Not in data/highlights.json: {', '.join(sorted(unknown))}")
    moments = fetch_offchain(a.refresh and not only)
    links = build_links(rows, moments)
    man = fetch_meta(links, a.refresh, only)
    if a.discover:
        refs = discover_refs(rows, moments, man)
        p(f"Discover: {len(refs)} moment(s) in NHL Breakaway set listings aren't in the card list yet")
        if a.limit: refs = refs[:a.limit]
        if refs: add_refs(refs, "found in NHL Breakaway set listing (fetch --discover)")
        rows = load_cards(); man = jload(os.path.join(DATA, "manifest.json"), {}); links = jload(os.path.join(DATA, "links.json"), {})
    # attach off-chain moment pointer to manifest
    by_uuid = {m["uuid"]: m for m in moments}
    for k, e in man.items():
        try:
            u = {t["name"]: t["value"] for t in jload(os.path.join(DATA, "meta", f"{k}.json")).get("attributes", [])}.get("Moment UUID")
        except Exception: u = None
        if u: e["moment_uuid"] = u
        if u in by_uuid: e["offchain_path"] = by_uuid[u]["file"]
    jdump(dict(sorted(man.items(), key=idkey)), os.path.join(DATA, "manifest.json"))
    write_cids(man)
    ok = sum(1 for e in man.values() if e.get("status") == 200)
    size = sum(os.path.getsize(os.path.join(DATA, "meta", f)) for f in os.listdir(os.path.join(DATA, "meta")))
    off = sum(1 for e in man.values() if e.get("offchain_path"))
    p(f"Fetch done: {len(rows)} highlights, {len(links)} linked, {ok} metadata OK, "
      f"{len(man) - ok} failed, {off} with off-chain data; data/meta = {size/1e6:.1f} MB"
      + (f" (refreshed ids: {', '.join(sorted(only, key=idkey))})" if only else ""))


# ------------------------------------------------------------------ offline model
def flat(o, pre="", out=None):
    out = {} if out is None else out
    if isinstance(o, dict):
        for k, v in o.items(): flat(v, f"{pre}.{k}" if pre else k, out)
    elif isinstance(o, list):
        if o and all(isinstance(x, dict) and "name" in x and "value" in x and len(x) <= 3 for x in o):
            for x in o: out[f"{pre}.{x['name']}"] = x["value"]
        elif o and all(isinstance(x, dict) and "purpose" in x for x in o):
            for x in o: flat({kk: vv for kk, vv in x.items() if kk != "purpose"}, f"{pre}.{x['purpose']}", out)
        elif all(not isinstance(x, (dict, list)) for x in o): out[pre] = ", ".join(map(str, o))
        else:
            for i, x in enumerate(o): flat(x, f"{pre}[{i}]", out)
    else: out[pre] = o
    return out


def load_model():
    rows = jload(os.path.join(DATA, "highlights.json"))
    if not rows: sys.exit("No data/highlights.json yet. Run: python breakaway.py fetch")
    man = jload(os.path.join(DATA, "manifest.json"), {})
    offc = {}
    sd = os.path.join(DATA, "offchain", "sets")
    if os.path.isdir(sd):
        for fn in os.listdir(sd):
            for m in jload(os.path.join(sd, fn), []): offc[m["uuid"]] = m
    for r in rows:
        k = str(r["id"]); e = man.get(k, {}); r["_man"] = e
        mp = os.path.join(DATA, "meta", f"{k}.json")
        r["_meta"] = jload(mp) if os.path.exists(mp) else None
        u = e.get("moment_uuid")
        if not u and r["_meta"]:
            u = {t["name"]: t["value"] for t in r["_meta"].get("attributes", [])}.get("Moment UUID")
        r["_off"] = offc.get(u)
    return rows


def summary(r):
    d = r["_meta"] or {}; a = {t["name"]: t["value"] for t in d.get("attributes", [])}
    off = r["_off"] or {}; oa = off.get("attributes") or {}; mo = oa.get("moment") or {}; g = oa.get("game") or {}
    pl = oa.get("primary_player") or {}; imgs = d.get("images") or {}; media = (d.get("assets") or {}).get("media") or []
    vid = next((m.get("uri") for m in media if m.get("purpose") == "primary-asset"), "") or next(
        (m.get("uri") for m in (off.get("assets") or {}).get("media") or [] if m.get("purpose") == "video"), "")
    front = imgs.get("front") or {}
    ln = r["_man"]
    return {
        "player": r["player"], "team": a.get("Team") or r["team"], "position": pl.get("position", ""),
        "jersey": pl.get("jersey", ""), "date": a.get("Date") or mo.get("date", ""), "play": a.get("Play") or mo.get("event", ""),
        "tagline": oa.get("collectible_tagline", ""), "period": mo.get("period", ""), "clock": mo.get("clock", ""),
        "final_score": g.get("final_score", ""), "venue": (g.get("venue") or {}).get("name", ""),
        "flags": ", ".join(k.replace("_", " ") for k, v in (mo.get("flags") or {}).items() if v),
        "set": a.get("Set") or r["set"], "list_set": r["set"], "series": a.get("Series") or (off.get("token_series") or {}).get("name", ""),
        "rarity": a.get("Rarity") or r["rarity"], "editions": int(a["Total Editions"]) if str(a.get("Total Editions", "")).isdigit() else (off.get("editions") or r["editions"] or ""),
        "description": re.sub(r"\n+Officially Licensed Product.*$", "", (d.get("description") or off.get("description") or "").strip(), flags=re.S).strip(),
        "issued": (d.get("issued") or off.get("issued") or "")[:10], "audio": a.get("Audio Type", ""),
        "lowest_ask": r.get("lowest_ask", ""), "fmv": r.get("fmv", ""),
        "cover_url": front.get("url", ""), "cover_ipfs": (front.get("altloc") or {}).get("ipfs", ""), "cover_sha256": front.get("sha256", ""),
        "video_url": vid, "breakaway_url": ln.get("source_url", "").replace("nft.nhlbreakaway.com/series", "nhlbreakaway.com/series"),
        "meta_url": ln.get("source_url", ""), "id": r["id"], "on_site": "yes" if r.get("on_site", True) else "no",
        "site_url": r.get("site_url", ""),
        "moment_uuid": ln.get("moment_uuid", ""), "raw_meta_path": ln.get("meta_path", ""),
        "offchain_path": ln.get("offchain_path", ""), "link_via": ln.get("link_via", ""),
        "source": " + ".join(x for x, ok in (("card list", True), ("nft. metadata", bool(r["_meta"])), ("Breakaway off-chain", bool(r["_off"]))) if ok),
        "missing": "" if r["_meta"] else ("no Breakaway token found" if not ln else f"metadata fetch failed ({ln.get('status')})"),
    }


def sort_rows(rows):
    for r in rows: r["_s"] = summary(r)
    rows.sort(key=lambda r: (r["_s"]["issued"] or "9999", r["_s"]["set"].lower(), r["_s"]["date"] or "", r["player"].lower(), idkey(r["id"])))
    for i, r in enumerate(rows, 1): r["_s"]["row"] = i
    return rows


def safe(s): return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:50]
def cover_path(r):
    for ext in (".png", ".jpg", ".jpeg", ".webp"):
        fp = os.path.join(OUT, "covers_full", f"{r['id']}{ext}")
        if os.path.exists(fp): return fp


COLS = [("row", "#"), ("cover", "Cover"), ("player", "Player"), ("team", "Team"), ("position", "Pos"), ("jersey", "Jersey"),
        ("date", "Highlight Date"), ("play", "Play"), ("tagline", "Highlight"), ("period", "Period"), ("clock", "Clock"),
        ("final_score", "Final Score"), ("venue", "Venue"), ("flags", "Flags"), ("set", "Set"), ("list_set", "Set (list label)"),
        ("series", "Series"), ("rarity", "Rarity"), ("editions", "Edition Size"), ("description", "Description"),
        ("issued", "Released"), ("audio", "Audio"), ("lowest_ask", "Lowest Ask"), ("fmv", "FMV"),
        ("cover_url", "Cover Art (full res)"), ("cover_ipfs", "Cover IPFS CID"), ("cover_sha256", "Cover sha256"),
        ("video_url", "Video"), ("breakaway_url", "NHL Breakaway Page"), ("meta_url", "Metadata URL"),
        ("cover_file", "Cover File"), ("id", "Card ID"), ("on_site", "In Original List"), ("site_url", "Card List Page"),
        ("moment_uuid", "Moment UUID"), ("raw_meta_path", "raw_meta_path"), ("offchain_path", "offchain_path"),
        ("link_via", "How Linked"), ("source", "Data Source"), ("missing", "Missing / Notes")]
LINKS = {"cover_url": "Cover art", "video_url": "Watch", "breakaway_url": "Breakaway", "meta_url": "JSON", "site_url": "The Fifth Line"}


def cmd_build(a):
    """Offline. Reads data/ (and output/covers_full/ for thumbnails if present)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XImg
    rows = sort_rows(load_model())
    try: from PIL import Image
    except ImportError: Image = None
    os.makedirs(OUT, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Highlights"
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    link_font = Font(color="0563C1", underline="single")

    def header(sheet, names):
        sheet.append(names)
        for c in sheet[1]: c.font, c.fill = hdr_font, hdr_fill
    header(ws, [h for _, h in COLS]); ws.freeze_panes = "D2"
    keys = [k for k, _ in COLS]; cc = keys.index("cover") + 1
    widths = {i: len(h) + 2 for i, (_, h) in enumerate(COLS, 1)}
    tdir = os.path.join(CACHE, "thumbs"); os.makedirs(tdir, exist_ok=True)
    for r in rows:
        s = r["_s"]; cp = cover_path(r)
        s["cover_file"] = f"covers/{s['row']:04d}_{safe(r['player'])}_{safe(s['set'])}{os.path.splitext(cp)[1]}" if cp else ""
        ws.append([("" if k == "cover" else s.get(k, "")) for k in keys]); ri = ws.max_row
        for ci, k in enumerate(keys, 1):
            c = ws.cell(ri, ci); c.alignment = Alignment(vertical="center"); v = c.value
            if k in LINKS and v: c.hyperlink = v; c.value = v = LINKS[k]; c.font = link_font
            elif k == "cover_file" and v: c.hyperlink = v; c.font = link_font
            widths[ci] = min(max(widths[ci], len(str(v if v is not None else "")) + 2), 70 if k == "description" else 48)
        if cp and Image:
            tp = os.path.join(tdir, f"{r['id']}.png")
            if not os.path.exists(tp) or os.path.getmtime(tp) < os.path.getmtime(cp):
                im = Image.open(cp).convert("RGB"); im.thumbnail((72, 72)); im.save(tp, "PNG")
            xi = XImg(tp); xi.anchor = f"{get_column_letter(cc)}{ri}"; ws.add_image(xi); ws.row_dimensions[ri].height = 58
    widths[cc] = 11
    for ci, w in widths.items(): ws.column_dimensions[get_column_letter(ci)].width = max(w, 6)
    ws.auto_filter.ref = ws.dimensions
    # every field: on-chain (meta.*) + off-chain (offchain.*), flattened
    fl = []
    for r in rows:
        f = {"#": r["_s"]["row"], "id": r["id"], "raw_meta_path": r["_s"]["raw_meta_path"], "offchain_path": r["_s"]["offchain_path"]}
        if r["_meta"]: f.update({f"meta.{k}": v for k, v in flat(r["_meta"]).items()})
        if r["_off"]: f.update({f"offchain.{k}": v for k, v in flat(r["_off"]).items()})
        fl.append(f)
    cols = list(dict.fromkeys(k for f in fl for k in f))
    fs = wb.create_sheet("All Metadata"); header(fs, cols); fs.freeze_panes = "C2"
    for f in fl:
        fs.append([(json.dumps(v) if isinstance(v, (dict, list)) else (str(v)[:32000] if isinstance(v, str) else v)) for v in (f.get(c, "") for c in cols)])
    for i, c in enumerate(cols, 1): fs.column_dimensions[get_column_letter(i)].width = min(max(len(c), 10) + 2, 40)
    # IPFS sheet: every CID + sha256
    ip = wb.create_sheet("IPFS"); header(ip, ["#", "id", "field", "ipfs", "sha256", "url"])
    for r in rows:
        for e in (r["_man"].get("ipfs") or []): ip.append([r["_s"]["row"], r["id"], e["path"], e["ipfs"], e["sha256"], e["url"]])
    for col, w in zip("ABCDEF", (6, 8, 40, 56, 48, 80)): ip.column_dimensions[col].width = w
    bs = wb.create_sheet("By Set"); header(bs, ["Set (list label)", "Highlights", "With Breakaway metadata"])
    agg = {}
    for r in rows: x = agg.setdefault(r["set"], [0, 0]); x[0] += 1; x[1] += 1 if r["_meta"] else 0
    for k in sorted(agg): bs.append([k, *agg[k]])
    bs.column_dimensions["A"].width = 48; bs.column_dimensions["C"].width = 24
    wb.save(os.path.join(OUT, "highlights.xlsx"))
    jdump([{k: v for k, v in r["_s"].items()} for r in rows], os.path.join(OUT, "highlights_summary.json"))
    n = sum(1 for r in rows if r["_meta"]); o = sum(1 for r in rows if r["_off"])
    p(f"Built output/highlights.xlsx: {len(rows)} rows, {n} with nft. metadata, {o} with off-chain data, "
      f"{sum(1 for r in rows if cover_path(r))} with cover thumbnails, {len(cols)} flattened fields")
    return rows


# ------------------------------------------------------------------ covers
def cmd_covers(a):
    rows = load_model(); cd = os.path.join(OUT, "covers_full"); os.makedirs(cd, exist_ok=True)
    stats = {"ok": 0, "cached": 0, "fail": 0, "sha_mismatch": 0}; lock = threading.Lock()
    def one(r):
        d = r["_meta"] or {}; front = (d.get("images") or {}).get("front") or {}; url = front.get("url")
        if not url: return
        ext = os.path.splitext(url.split("?")[0])[1] or ".png"; fp = os.path.join(cd, f"{r['id']}{ext}")
        if os.path.exists(fp): 
            with lock: stats["cached"] += 1
            return
        cid, _ = asset_ipfs(d, "front")
        b, used = get_with_fallback(url, cid)
        with lock:
            if not b: stats["fail"] += 1; return
            if front.get("sha256") and b64sha(b).rstrip("=") != front["sha256"].rstrip("="): stats["sha_mismatch"] += 1
            wbytes(fp, b); stats["ok"] += 1
    with cf.ThreadPoolExecutor(THR.sem._value) as ex:
        for n, _ in enumerate(ex.map(one, rows), 1):
            if n % 250 == 0: p(f"  covers {n}/{len(rows)}")
    size = sum(os.path.getsize(os.path.join(cd, f)) for f in os.listdir(cd))
    p(f"Covers: {stats}; output/covers_full = {size/1e6:.0f} MB")


# ------------------------------------------------------------------ package for sharing
FRIEND_README = """NHL BREAKAWAY HIGHLIGHTS
Unofficial fan preservation project. Not affiliated with or endorsed by the NHL, NHLPA, NHLAA, Sweet (SocialSweet Inc.) or any team. All trademarks, names, logos, videos and images belong to their respective owners. This repository contains no media; anyone who downloads media with the tool does so only for personal archival use. Takedown requests are honored: please open a GitHub issue.
Generated {when}. {n} highlights; {m} with full NHL Breakaway details, {c} with cover art.

FILES
  highlights.xlsx   Sheet "Highlights": one row per highlight with a small cover thumbnail (column B).
                    Sheet "All Metadata": every field NHL Breakaway publishes for each highlight.
                    Sheet "IPFS": IPFS copies (CID + sha256) of every image/video, in case the site goes away.
                    Sheet "By Set": highlight count per set.
  covers/           Full-resolution cover art named <row #>_<player>_<set>, matching the "#" column.
  breakaway.py      The script that made all this (optional; see below).

KEY COLUMNS
  # / Cover / Player / Team / Pos / Jersey     who and the cover art
  Highlight Date / Play / Highlight / Period / Clock / Final Score / Venue / Flags   what happened
  Set / Series / Rarity / Edition Size / Released   collectible details (official NHL Breakaway values;
                    "Set (list label)" is the set name as written in data/highlights.json)
  Description       The official write-up, good for album captions
  Lowest Ask / FMV  Price snapshot (USD) saved with the card list in October 2026
  Cover Art, Video, NHL Breakaway Page, Metadata URL, Card List Page   clickable links
  Cover IPFS CID / Cover sha256   permanent IPFS address and checksum of the cover image
  Cover File        matching image in covers/
  Data Source / Missing / Notes   where the row's data came from, or why something is missing

RE-RUNNING OR DOWNLOADING VIDEOS (optional)
  1. Install Python 3 from https://www.python.org/downloads/ (tick "Add Python to PATH" on Windows).
  2. In this folder open a terminal and run:   pip install -r requirements.txt
     (or: pip install openpyxl pillow)
  3. Rebuild everything:                       python breakaway.py all
  4. Videos are OFF by default. To save them:  python breakaway.py videos --out-dir videos
     It shows the estimated total size and asks y/n first. Try --limit 5 for a small test,
     --series "Series 3" or --ids 12,40 to pick some. Re-running resumes and skips finished files.
  5. Printable flip-book PDF:                  python breakaway.py album   (saved as output/album.pdf)
     Options: --layout 1up|2x2|3x3, --page letter|a4, --set/--series/--player/--team "name", --bleed
"""


def cmd_package(a):
    rows = sort_rows(load_model())
    cov = os.path.join(OUT, "covers"); os.makedirs(cov, exist_ok=True)
    for fn in os.listdir(cov): os.remove(os.path.join(cov, fn))
    names = []
    for r in rows:
        cp = cover_path(r)
        if cp:
            fn = f"{r['_s']['row']:04d}_{safe(r['player'])}_{safe(r['_s']['set'])}{os.path.splitext(cp)[1]}"
            os.link(cp, os.path.join(cov, fn)); names.append(fn)
    with zipfile.ZipFile(os.path.join(OUT, "covers.zip"), "w", zipfile.ZIP_STORED) as z:
        for fn in names: z.write(os.path.join(cov, fn), f"covers/{fn}")
    readme = FRIEND_README.format(when=time.strftime("%Y-%m-%d"), n=len(rows), m=sum(1 for r in rows if r["_meta"]), c=len(names))
    open(os.path.join(OUT, "README.txt"), "w").write(readme)
    zp = os.path.join(ROOT, "Breakaway-Highlights.zip")
    with zipfile.ZipFile(zp, "w") as z:
        D = zipfile.ZIP_DEFLATED
        z.write(os.path.join(OUT, "highlights.xlsx"), "Breakaway-Highlights/highlights.xlsx", D)
        z.writestr("Breakaway-Highlights/README.txt", readme, D)
        z.write(os.path.abspath(__file__), "Breakaway-Highlights/breakaway.py", D)
        z.write(os.path.join(ROOT, "requirements.txt"), "Breakaway-Highlights/requirements.txt", D)
        for fn in names: z.write(os.path.join(cov, fn), f"Breakaway-Highlights/covers/{fn}", zipfile.ZIP_STORED)
    p(f"Packaged {zp} ({os.path.getsize(zp)/1e6:.0f} MB, {len(names)} covers)")


# ------------------------------------------------------------------ videos (opt-in)
def video_size(url):
    """Size from a 1-byte ranged GET (Content-Range total); falls back to HEAD."""
    st, _, h = http(url, timeout=30, headers={"Range": "bytes=0-0"})
    m = re.search(r"/(\d+)$", (h or {}).get("Content-Range", "") or "")
    if m: return int(m.group(1))
    st, _, h = http(url, method="HEAD", timeout=30)
    return int(h["Content-Length"]) if h and h.get("Content-Length") else None


def cmd_videos(a):
    rows = sort_rows(load_model())
    if a.series: rows = [r for r in rows if r["_s"]["series"].lower() == a.series.lower()]
    if a.ids: want = {x.strip() for x in a.ids.split(",") if x.strip()}; rows = [r for r in rows if str(r["id"]) in want]
    rows = [r for r in rows if r["_s"]["video_url"]]
    if a.limit: rows = rows[: a.limit]
    out = os.path.abspath(a.out_dir); os.makedirs(out, exist_ok=True)
    def fname(r): return os.path.join(out, video_name(r))
    todo = [r for r in rows if not os.path.exists(fname(r))]
    p(f"{len(rows)} videos selected, {len(rows) - len(todo)} already saved, {len(todo)} to download into {out}")
    if not todo: return
    sample = todo if len(todo) <= 30 else todo[:: max(1, len(todo) // 30)][:30]
    sizes = []
    for r in sample:
        cid, _ = asset_ipfs(r["_meta"], "primary-asset")
        n = next((x for x in (video_size(u) for u in candidate_urls(r["_s"]["video_url"], cid)[:2]) if x), None)
        if n: sizes.append(n)
    est = (sum(sizes) / len(sizes) * len(todo)) if sizes else None
    p(f"Estimated total: {est/1e9:.2f} GB (avg {sum(sizes)/len(sizes)/1e6:.1f} MB from {len(sizes)} size checks)" if est else "Could not estimate size.")
    if not a.yes and input("Download now? [y/N] ").strip().lower() not in ("y", "yes"): p("Cancelled."); return
    def one(r):
        dst = fname(r); part = dst + ".part"; have = os.path.getsize(part) if os.path.exists(part) else 0
        cid, _ = asset_ipfs(r["_meta"], "primary-asset")
        urls = candidate_urls(r["_s"]["video_url"], cid)
        for i in range(6):
            src = urls[min(i // 2, len(urls) - 1)]  # 2 tries per source, then the next gateway
            THR.sem.acquire(); time.sleep(THR.delay)
            try:
                hdr = {**HDR, **({"Range": f"bytes={have}-"} if have else {})}
                with urllib.request.urlopen(urllib.request.Request(src, headers=hdr), timeout=120) as resp:
                    mode = "ab" if have and resp.status == 206 else "wb"
                    with open(part, mode) as f:
                        while True:
                            ch = resp.read(1 << 20)
                            if not ch: break
                            f.write(ch)
                os.replace(part, dst); THR.ok(); return True
            except urllib.error.HTTPError as e:
                if e.code in (401, 403, 404) and i // 2 >= len(urls) - 1: p(f"  skip {r['id']}: HTTP {e.code}"); return False
                THR.bad(); time.sleep(min(60, 2 ** (i + 1)))
            except Exception:
                THR.bad(); time.sleep(min(60, 2 ** (i + 1)))
                have = os.path.getsize(part) if os.path.exists(part) else 0
            finally: THR.sem.release()
        return False
    with cf.ThreadPoolExecutor(THR.sem._value) as ex:
        ok = sum(1 for x in ex.map(one, todo) if x)
    p(f"Videos: {ok}/{len(todo)} downloaded to {out}")


# ------------------------------------------------------------------ IPFS / gateways / preservation
DEFAULT_GATEWAYS = ["https://ipfs.io/ipfs/", "https://dweb.link/ipfs/", "https://w3s.link/ipfs/",
                    "https://gateway.pinata.cloud/ipfs/", "https://4everland.io/ipfs/"]
GATEWAYS = list(DEFAULT_GATEWAYS)
IPFS_ONLY = False


def write_cids(man):
    """data/cids.txt: one CID per line (deduplicated) so anyone can `ipfs pin add` them."""
    seen = {}
    for k, e in sorted(man.items(), key=idkey):
        for x in e.get("ipfs") or []:
            cid = x["ipfs"].replace("ipfs://", "").strip("/")
            seen.setdefault(cid, f"{k} {x['path']}")
    with open(os.path.join(DATA, "cids.txt"), "w") as f:
        f.write("# IPFS CIDs referenced by the NHL Breakaway metadata in data/meta/ (deduplicated)\n")
        f.write("# pin all:  grep -v '^#' data/cids.txt | cut -d' ' -f1 | xargs -n 50 ipfs pin add\n")
        for cid, where in seen.items(): f.write(f"{cid}  # {where}\n")
    return len(seen)


def asset_ipfs(meta, purpose_or_path):
    """IPFS CID for a media purpose ('primary-asset', 'card-image', ...) or image key ('front')."""
    d = meta or {}
    if purpose_or_path in (d.get("images") or {}):
        x = d["images"][purpose_or_path]; return (x.get("altloc") or {}).get("ipfs"), x.get("sha256")
    for m in (d.get("assets") or {}).get("media") or []:
        if m.get("purpose") == purpose_or_path: return (m.get("altloc") or {}).get("ipfs"), m.get("sha256")
    return None, None


def candidate_urls(primary, cid):
    urls = [] if IPFS_ONLY or not primary else [primary]
    if cid:
        c = cid.replace("ipfs://", "").strip("/")
        urls += [g.rstrip("/") + "/" + c for g in GATEWAYS]
    return urls


def get_with_fallback(primary, cid, timeout=120):
    for u in candidate_urls(primary, cid):
        st, b, _ = http(u, timeout=timeout, tries=3)
        if b: return b, u
    return None, None


def sha_ok(b, expected):
    return (not expected) or b64sha(b).rstrip("=") == expected.rstrip("=")


def cmd_verify(a):
    rows = load_model(); man = jload(os.path.join(DATA, "manifest.json"), {})
    res = {"meta_ok": 0, "meta_bad": 0, "meta_missing": 0, "cover_ok": 0, "cover_bad": 0, "video_ok": 0, "video_bad": 0}
    bad = []
    for k, e in sorted(man.items(), key=idkey):
        fp = os.path.join(DATA, "meta", f"{k}.json")
        if not os.path.exists(fp): res["meta_missing"] += 1; continue
        if hashlib.sha256(open(fp, "rb").read()).hexdigest() == e.get("sha256"): res["meta_ok"] += 1
        else: res["meta_bad"] += 1; bad.append(f"meta {k}")
    vdir = os.path.abspath(a.out_dir)
    for r in sort_rows(rows):
        cp = cover_path(r)
        if cp:
            _, sh = asset_ipfs(r["_meta"], "front")
            if sha_ok(open(cp, "rb").read(), sh): res["cover_ok"] += 1
            else: res["cover_bad"] += 1; bad.append(f"cover {r['id']}")
        vp = os.path.join(vdir, video_name(r))
        if os.path.exists(vp):
            _, sh = asset_ipfs(r["_meta"], "primary-asset")
            h = hashlib.sha256()
            with open(vp, "rb") as f:
                for ch in iter(lambda: f.read(1 << 20), b""): h.update(ch)
            ok = (not sh) or base64.urlsafe_b64encode(h.digest()).decode().rstrip("=") == sh.rstrip("=")
            res["video_ok" if ok else "video_bad"] += 1
            if not ok: bad.append(f"video {r['id']}")
    jdump({"checked_at": now(), "results": res, "mismatches": bad}, os.path.join(OUT, "verify_report.json"))
    p(f"Verify: {res}")
    for b in bad[:50]: p("  mismatch:", b)
    p("Note: a cover/video mismatch means the CDN now serves different bytes than the sha256 recorded in the "
      "metadata; the IPFS copy (same CID) is the original.")


# ------------------------------------------------------------------ add moments by URL
URL_RE = re.compile(r"(?:nft\.)?nhlbreakaway\.com/series/([A-Za-z0-9]{6,12})/(\d{6,})|^([A-Za-z0-9]{6,12})/(\d{6,})$")


def parse_moment_refs(items):
    out = []
    for it in items:
        it = it.strip()
        if not it or it.startswith("#"): continue
        m = URL_RE.search(it)
        if m: out.append((m.group(1) or m.group(3), m.group(2) or m.group(4)))
        else: p("  can't read:", it)
    return out


def cmd_add(a):
    items = list(a.refs or [])
    extra = os.path.join(DATA, "extra_moments.txt")
    if os.path.exists(extra): items += open(extra, encoding="utf-8").read().splitlines()
    refs = parse_moment_refs(items)
    if not refs: p("Nothing to add. Pass URLs or put them in data/extra_moments.txt"); return
    add_refs(refs, "added by URL")


def add_refs(refs, via):
    rows = jload(os.path.join(DATA, "highlights.json"), []); links = jload(os.path.join(DATA, "links.json"), {})
    man = jload(os.path.join(DATA, "manifest.json"), {})
    known_uuid = {e.get("moment_uuid"): k for k, e in man.items() if e.get("moment_uuid")}
    known_tok = {(v["series_id"], str(int(v["token_id"]) >> 48)): k for k, v in links.items()}
    added = 0
    nid = max([int(r["id"]) for r in rows if str(r["id"]).isdigit()] + [int(k) for k in man if k.isdigit()] + [0]) + 1
    for sid, tok in refs:
        if (sid, str(int(tok) >> 48)) in known_tok: p(f"  already have {sid}/{tok} (id {known_tok[(sid, str(int(tok) >> 48))]})"); continue
        st, b, h, tok = get_nft(sid, tok); url = f"{NFT}/{sid}/{tok}"
        if not b: p(f"  {url}: HTTP {st}, skipped"); continue
        d = json.loads(b); at = {t["name"]: t["value"] for t in d.get("attributes", [])}
        u = at.get("Moment UUID")
        if u in known_uuid: p(f"  {sid}/{tok} is the same moment as id {known_uuid[u]}, skipped"); continue
        k = str(nid); nid += 1  # new card ids go after the existing ones; nothing is renumbered
        wbytes(os.path.join(DATA, "meta", f"{k}.json"), b)
        man[k] = {"id": int(k), "source_url": url, "fetched_at": now(), "status": st, "link_via": via,
                  "meta_path": f"data/meta/{k}.json", "bytes": len(b), "sha256": hashlib.sha256(b).hexdigest(),
                  "content_type": h.get("Content-Type") if h else None, "ipfs": ipfs_entries(d), "moment_uuid": u}
        links[k] = {"series_id": sid, "token_id": tok, "exact": True, "via": via}
        if at.get("Set UUID"):  # off-chain set listing for this moment
            sp = os.path.join(DATA, "offchain", "sets", f"{at['Set UUID']}.json")
            if not os.path.exists(sp):
                st2, b2, _ = http(f"{APIGW}/catalog/search/moment/set/{at['Set UUID']}?campaign_id=826")
                if b2: wbytes(sp, b2)
            if os.path.exists(sp) and any(m.get("uuid") == u for m in jload(sp, [])):
                man[k]["offchain_path"] = f"data/offchain/sets/{at['Set UUID']}.json"
        rows.append({"id": int(k), "player": at.get("Player", d.get("name", "")), "set": at.get("Set", ""),
                     "rarity": at.get("Rarity", ""), "editions": int(at["Total Editions"]) if str(at.get("Total Editions", "")).isdigit() else None,
                     "team": at.get("Team", ""), "lowest_ask": "", "fmv": "", "on_site": False})
        known_uuid[u] = k; known_tok[(sid, str(int(tok) >> 48))] = k; added += 1; p(f"  added {k}: {at.get('Player')} / {at.get('Set')}")
    rows.sort(key=lambda r: idkey(r["id"]))
    jdump(rows, os.path.join(DATA, "highlights.json"))
    jdump(dict(sorted(links.items(), key=idkey)), os.path.join(DATA, "links.json"))
    jdump(dict(sorted(man.items(), key=idkey)), os.path.join(DATA, "manifest.json"))
    write_cids(man); p(f"Added {added} moment(s). Run `python breakaway.py build` to refresh the spreadsheet.")


# ------------------------------------------------------------------ offline viewer
VIEWER = r"""<!doctype html><html><head><meta charset="utf-8"><title>breakaway-data viewer</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
body{margin:0;font:14px system-ui,sans-serif;background:#0f141a;color:#e6edf3}
header{position:sticky;top:0;background:#16202b;padding:10px;display:flex;flex-wrap:wrap;gap:8px;z-index:2}
input,select{background:#0f141a;color:#e6edf3;border:1px solid #33414f;border-radius:6px;padding:6px}
#count{align-self:center;color:#8b9aa8}
#grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:10px;padding:10px}
.card{background:#16202b;border-radius:8px;overflow:hidden;cursor:pointer}.card:hover{outline:2px solid #3b82f6}
.card img{width:100%;aspect-ratio:1;object-fit:cover;background:#222;display:block}
.card div{padding:6px 8px}.card b{display:block}.card small{color:#8b9aa8}
#dlg{position:fixed;inset:0;background:#000c;display:none;overflow:auto;z-index:3}
#box{max-width:1000px;margin:30px auto;background:#16202b;border-radius:10px;padding:16px}
#box img,#box video{max-width:100%;max-height:60vh;display:block;margin:0 auto 10px}
table{border-collapse:collapse;width:100%;font-size:12px}td{border-top:1px solid #26323e;padding:3px 6px;vertical-align:top;word-break:break-all}
td:first-child{color:#8b9aa8;white-space:nowrap;word-break:normal}a{color:#60a5fa}
#more{display:block;margin:10px auto;padding:8px 20px}
</style></head><body>
<header><input id="q" placeholder="Search player, set, description…" size="30">
<select id="f_team"></select><select id="f_set"></select><select id="f_rarity"></select><select id="f_series"></select>
<span id="count"></span></header>
<div id="grid"></div><button id="more">Show more</button>
<footer style="padding:16px;color:#8b9aa8;font-size:12px;text-align:center">Unofficial fan preservation project. Not affiliated with or endorsed by the NHL, NHLPA, NHLAA, Sweet (SocialSweet Inc.) or any team. All trademarks, names, logos, videos and images belong to their respective owners. This repository contains no media; anyone who downloads media with the tool does so only for personal archival use. Takedown requests are honored: please open a GitHub issue at https://github.com/breakaway-vault/breakaway-data/issues.</footer>
<div id="dlg"><div id="box"></div></div>
<script>
const DATA = __DATA__;
const GW = __GATEWAYS__;
const $ = s => document.querySelector(s);
const F = {team:'team', set:'set', rarity:'rarity', series:'series'};
function opts(sel, key, label){const v=[...new Set(DATA.map(r=>r.s[key]).filter(Boolean))].sort();
  sel.innerHTML=`<option value="">All ${label}</option>`+v.map(x=>`<option>${esc(x)}</option>`).join('');}
function esc(s){return String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]))}
opts($('#f_team'),'team','teams');opts($('#f_set'),'set','sets');opts($('#f_rarity'),'rarity','rarities');opts($('#f_series'),'series','series');
function ipfsUrl(cid){return cid?GW[0]+cid.replace('ipfs://',''):''}
function imgTag(r,cls){const s=r.s;const alts=[r.local_cover,s.cover_url,ipfsUrl(s.cover_ipfs)].filter(Boolean);
  return `<img loading="lazy" src="${esc(alts[0])}" data-alts="${esc(alts.slice(1).join('|'))}" onerror="nextSrc(this)" alt="">`}
function nextSrc(img){const a=(img.dataset.alts||'').split('|').filter(Boolean);if(!a.length){img.onerror=null;return}
  img.src=a.shift();img.dataset.alts=a.join('|')}
let list=[],shown=0;
function apply(){const q=$('#q').value.toLowerCase();
  const fv={team:$('#f_team').value,set:$('#f_set').value,rarity:$('#f_rarity').value,series:$('#f_series').value};
  list=DATA.filter(r=>(!q||r.text.includes(q))&&Object.entries(fv).every(([k,v])=>!v||r.s[k]===v));
  shown=0;$('#grid').innerHTML='';more();}
function more(){const part=list.slice(shown,shown+120);shown+=part.length;
  $('#grid').insertAdjacentHTML('beforeend',part.map(r=>`<div class="card" data-i="${r.i}">${imgTag(r)}<div><b>${esc(r.s.player)}</b><small>${esc(r.s.set)} · ${esc(r.s.rarity)} · ${esc(r.s.date)}</small></div></div>`).join(''));
  $('#count').textContent=`${list.length} of ${DATA.length} highlights`;$('#more').style.display=shown<list.length?'':'none';}
function rowsOf(o,pre=''){let out=[];if(o&&typeof o==='object'){for(const[k,v]of Object.entries(o)){out=out.concat(rowsOf(v,pre?pre+'.'+k:k))}}else out.push([pre,o]);return out}
function link(v){v=String(v??'');if(/^https?:\/\//.test(v))return `<a href="${esc(v)}" target="_blank">${esc(v)}</a>`;
  if(/^ipfs:\/\//.test(v))return `<a href="${esc(ipfsUrl(v))}" target="_blank">${esc(v)}</a>`;return esc(v)}
function open_(i){const r=DATA[i],s=r.s;
  const vids=[r.local_video,s.video_url,ipfsUrl(s.video_ipfs)].filter(Boolean);
  const video=r.local_video?`<video controls preload="metadata" src="${esc(r.local_video)}"></video>`:
    `<p>Video: ${vids.map(v=>`<a href="${esc(v)}" target="_blank">${v.startsWith('http')&&v.includes('/ipfs/')?'IPFS copy':'stream / download'}</a>`).join(' · ')||'none'}</p>`;
  const tbl=(title,o)=>`<h3>${title}</h3><table>${rowsOf(o).map(([k,v])=>`<tr><td>${esc(k)}</td><td>${link(v)}</td></tr>`).join('')}</table>`;
  $('#box').innerHTML=`<button onclick="$('#dlg').style.display='none'" style="float:right">Close</button>
  <h2>${esc(s.player)} — ${esc(s.set)}</h2>${imgTag(r)}${video}<p>${esc(s.description)}</p>
  ${tbl('Summary',s)}${tbl('IPFS copies (CID / sha256)',r.ipfs)}${tbl('Source',r.src)}${tbl('All metadata (nft.nhlbreakaway.com)',r.meta||{})}${tbl('Off-chain (NHL Breakaway set listing)',r.off||{})}`;
  $('#dlg').style.display='block';$('#dlg').scrollTop=0;}
$('#grid').onclick=e=>{const c=e.target.closest('.card');if(c)open_(+c.dataset.i)};
$('#dlg').onclick=e=>{if(e.target.id==='dlg')$('#dlg').style.display='none'};
document.onkeydown=e=>{if(e.key==='Escape')$('#dlg').style.display='none'};
['#q','#f_team','#f_set','#f_rarity','#f_series'].forEach(s=>$(s).oninput=apply);$('#more').onclick=more;apply();
</script></body></html>"""


def video_name(r): return f"{r['_s']['row']:04d}_{safe(r['player'])}_{safe(r['_s']['set'])}.mp4"


def cmd_view(a):
    rows = sort_rows(load_model()); vd = os.path.join(ROOT, "viewer"); os.makedirs(vd, exist_ok=True)
    vdir = os.path.abspath(a.out_dir); data = []
    for i, r in enumerate(rows):
        s = dict(r["_s"]); s["video_ipfs"], _ = asset_ipfs(r["_meta"], "primary-asset")
        cp = cover_path(r); vp = os.path.join(vdir, video_name(r))
        rel = lambda x: os.path.relpath(x, vd).replace(os.sep, "/")
        data.append({"i": i, "s": s, "meta": r["_meta"], "off": {k: v for k, v in (r["_off"] or {}).items() if k not in ("series", "token_set", "badges")},
                     "ipfs": {e["path"]: f"{e['ipfs']}  sha256={e['sha256']}" for e in (r["_man"].get("ipfs") or [])},
                     "src": {k: r["_man"].get(k) for k in ("source_url", "fetched_at", "sha256", "meta_path", "offchain_path", "link_via")},
                     "local_cover": rel(cp) if cp else "", "local_video": rel(vp) if os.path.exists(vp) else "",
                     "text": " ".join(str(s.get(k, "")) for k in ("player", "team", "set", "series", "rarity", "description", "tagline", "date")).lower()})
    out = VIEWER.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/")).replace("__GATEWAYS__", json.dumps(GATEWAYS))
    open(os.path.join(vd, "index.html"), "w", encoding="utf-8").write(out)
    p(f"Wrote viewer/index.html ({len(out)/1e6:.1f} MB, {len(data)} highlights, "
      f"{sum(1 for d in data if d['local_cover'])} local covers, {sum(1 for d in data if d['local_video'])} local videos). Open it in a browser.")



# ------------------------------------------------------------------ printable album (PDF)
DISCLAIMER = ("Unofficial fan preservation project. Not affiliated with or endorsed by the NHL, NHLPA, NHLAA, "
              "Sweet (SocialSweet Inc.) or any team. All trademarks, names, logos, videos and images belong to "
              "their respective owners. For personal archival use only.")


def pdf_text(s):
    """Keep characters the built-in PDF fonts can draw (cp1252); transliterate the rest (e.g. c-caron -> c)."""
    out = []
    for ch in str(s if s is not None else ""):
        try: ch.encode("cp1252"); out.append(ch)
        except UnicodeEncodeError:
            out.append(unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode() or "?")
    return "".join(out)


def album_sections(rows):
    """[(series, set, [rows])]: every card of a set together, sets ordered by their first release date."""
    groups = {}
    for r in rows: groups.setdefault((r["_s"]["series"] or "", r["_s"]["set"]), []).append(r)
    secs = []
    for (series, name), rs in groups.items():
        rs.sort(key=lambda r: (r["_s"]["date"] or "", r["_s"]["row"]))
        secs.append([series, name, rs])
    secs.sort(key=lambda s: (min((r["_s"]["issued"] or "9999") for r in s[2]), s[0], s[1].lower()))
    return secs


def cmd_album(a):
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib.units import mm, inch
        from reportlab.lib.utils import ImageReader
    except ImportError:
        sys.exit("The album needs reportlab: python -m pip install -r requirements.txt")
    from PIL import Image
    rows = sort_rows(load_model())
    def has(val, *fields):
        return not val or any(val.lower() in str(f or "").lower() for f in fields)
    rows = [r for r in rows if has(a.set, r["_s"]["set"], r["_s"]["list_set"]) and has(a.series, r["_s"]["series"])
            and has(a.player, r["_s"]["player"]) and has(a.team, r["_s"]["team"], r["team"])]
    if not rows: sys.exit("No highlights match those filters.")
    if any(not cover_path(r) for r in rows):
        p("Some covers aren't downloaded yet; running `covers` first."); cmd_covers(a)
    secs = album_sections(rows)
    cols, rws = {"1up": (1, 1), "2x2": (2, 2), "3x3": (3, 3)}[a.layout]; per = cols * rws
    TW, TH = A4 if a.page == "a4" else letter            # trim size
    BL = 3 * mm if a.bleed else 0                         # bleed on each side
    M = 0.5 * inch                                        # safe margin inside trim
    # --- page plan: cover, TOC, then per section a title page + card pages
    toc_lines = []
    for s in secs: toc_lines += ([("series", s[0])] if not toc_lines or s[0] != prev_series else []) + [("set", s)]; prev_series = s[0]
    toc_per_page = int((TH - 2 * M - 70) // 15)
    toc_pages = max(1, -(-len(toc_lines) // toc_per_page))
    page = 1 + toc_pages + 1; start = {}
    for s in secs:
        start[id(s)] = page; page += 1 + -(-len(s[2]) // per)
    total_pages = page - 1
    out = os.path.abspath(a.out or os.path.join(OUT, "album.pdf")); os.makedirs(os.path.dirname(out), exist_ok=True)
    c = canvas.Canvas(out, pagesize=(TW + 2 * BL, TH + 2 * BL), pageCompression=1)
    c.setTitle("Breakaway Highlights Album"); c.setAuthor("breakaway-data (unofficial fan archive)")
    state = {"n": 0}

    def crop_marks():
        if not BL: return
        c.setLineWidth(0.25); c.setStrokeColorRGB(0, 0, 0); L = 6 * mm
        for x in (BL, BL + TW):
            for y in (BL, BL + TH):
                dx = -1 if x == BL else 1; dy = -1 if y == BL else 1
                c.line(x + dx * 1 * mm, y, x + dx * (1 * mm + L), y)   # horizontal mark outside trim
                c.line(x, y + dy * 1 * mm, x, y + dy * (1 * mm + L))   # vertical mark outside trim

    def begin():
        state["n"] += 1; c.translate(BL, BL)   # draw in trim coordinates

    def end(number=True):
        if number:
            c.setFont("Helvetica", 8); c.setFillGray(0.4)
            c.drawCentredString(TW / 2, M / 2, f"{state['n']}")
        c.translate(-BL, -BL); crop_marks(); c.showPage()

    def wrap(text, font, size, width):
        words, lines, cur = pdf_text(text).split(), [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if c.stringWidth(t, font, size) <= width: cur = t
            else: lines.append(cur); cur = w
        if cur: lines.append(cur)
        return lines

    def fit(text, font, size, width):
        t = pdf_text(text)
        while t and c.stringWidth(t, font, size) > width: t = t[:-2] + "…" if len(t) > 2 else ""
        return t

    # cover
    begin(); c.setFillGray(0.08); c.rect(-BL, -BL, TW + 2 * BL, TH + 2 * BL, stroke=0, fill=1)
    c.setFillGray(1); c.setFont("Helvetica-Bold", 34); c.drawCentredString(TW / 2, TH * 0.62, "Breakaway Highlights Album")
    c.setFont("Helvetica", 16)
    c.drawCentredString(TW / 2, TH * 0.62 - 40, f"{len(rows):,} highlights · {len(secs)} sets")
    first = min((r["_s"]["issued"] for r in rows if r["_s"]["issued"]), default=""); last = max((r["_s"]["issued"] for r in rows if r["_s"]["issued"]), default="")
    if first: c.drawCentredString(TW / 2, TH * 0.62 - 62, f"Released {first[:7]} to {last[:7]}")
    c.setFont("Helvetica", 7); c.setFillGray(0.75); y = M + 30
    for ln in reversed(wrap(DISCLAIMER, "Helvetica", 7, TW - 2 * M)): c.drawCentredString(TW / 2, y, ln); y += 9
    end(number=False)
    # table of contents
    for tp in range(toc_pages):
        begin(); c.setFillGray(0); c.setFont("Helvetica-Bold", 20)
        c.drawString(M, TH - M - 10, "Contents" + (" (continued)" if tp else ""))
        y = TH - M - 50
        for kind, v in toc_lines[tp * toc_per_page:(tp + 1) * toc_per_page]:
            if kind == "series":
                c.setFont("Helvetica-Bold", 11); c.setFillGray(0.2); c.drawString(M, y, pdf_text(v or "Other"))
            else:
                c.setFont("Helvetica", 10); c.setFillGray(0)
                label = fit(f"{v[1]}  ({len(v[2])})", "Helvetica", 10, TW - 2 * M - 60)
                c.drawString(M + 14, y, label); c.drawRightString(TW - M, y, str(start[id(v)]))
                lw = c.stringWidth(label, "Helvetica", 10); c.setDash(1, 2); c.setStrokeGray(0.7)
                c.line(M + 18 + lw, y + 2, TW - M - 24, y + 2); c.setDash()
            y -= 15
        end()
    # sections
    max_px = a.max_px
    def image_for(path):
        if not max_px and not a.jpeg: return ImageReader(path)
        im = Image.open(path).convert("RGB")
        if max_px: im.thumbnail((max_px, max_px))
        b = io.BytesIO(); im.save(b, "JPEG", quality=a.jpeg or 88, optimize=True); b.seek(0)
        return ImageReader(b)
    done = 0
    for s in secs:
        series, setname, rs = s
        # set title page
        begin(); c.setFillGray(0)
        c.setFont("Helvetica", 12); c.setFillGray(0.35); c.drawCentredString(TW / 2, TH * 0.66, pdf_text(series))
        c.setFillGray(0); c.setFont("Helvetica-Bold", 28)
        for i, ln in enumerate(wrap(setname, "Helvetica-Bold", 28, TW - 2 * M)): c.drawCentredString(TW / 2, TH * 0.6 - i * 34, ln)
        mix = {}
        for r in rs: mix[r["_s"]["rarity"] or "?"] = mix.get(r["_s"]["rarity"] or "?", 0) + 1
        c.setFont("Helvetica", 13)
        c.drawCentredString(TW / 2, TH * 0.47, f"{len(rs)} highlight{'s' if len(rs) != 1 else ''}")
        c.drawCentredString(TW / 2, TH * 0.47 - 20, pdf_text(" · ".join(f"{k} {v}" for k, v in sorted(mix.items(), key=lambda x: -x[1]))))
        rel = rs[0]["_s"]["issued"]
        if rel: c.setFillGray(0.35); c.drawCentredString(TW / 2, TH * 0.47 - 40, f"Released {rel}")
        desc = re.sub(r"<[^>]+>", " ", html.unescape(((rs[0]["_off"] or {}).get("token_set") or {}).get("description") or ""))
        if desc.strip():
            c.setFont("Helvetica-Oblique", 9); c.setFillGray(0.3); y = TH * 0.47 - 75
            for ln in wrap(desc, "Helvetica-Oblique", 9, TW - 3 * M)[:12]: c.drawCentredString(TW / 2, y, ln); y -= 12
        end()
        # card pages
        fs = {1: 12, 2: 9, 3: 7}[cols]; lh = fs * 1.3; textblock = 5 * lh + 6
        gw, gh = TW - 2 * M, TH - 2 * M - 22; cw, ch = gw / cols, gh / rws; pad = 8
        for i0 in range(0, len(rs), per):
            begin(); c.setFillGray(0.4); c.setFont("Helvetica", 8)
            c.drawString(M, TH - M - 8, fit(f"{series} · {setname}", "Helvetica", 8, TW - 2 * M))
            for j, r in enumerate(rs[i0:i0 + per]):
                s_ = r["_s"]; col, row = j % cols, j // cols
                x0 = M + col * cw + pad; y_top = TH - M - 22 - row * ch - pad; w = cw - 2 * pad; h = ch - 2 * pad
                ih = h - textblock; cp = cover_path(r)
                if cp:
                    iw0, ih0 = Image.open(cp).size; sc = min(w / iw0, ih / ih0); dw, dh = iw0 * sc, ih0 * sc
                    c.drawImage(image_for(cp), x0 + (w - dw) / 2, y_top - dh, dw, dh)
                    yt = y_top - dh - lh
                else:
                    c.setStrokeGray(0.7); c.rect(x0, y_top - ih, w, ih); yt = y_top - ih - lh
                c.setFillGray(0); c.setFont("Helvetica-Bold", fs + 1)
                c.drawString(x0, yt, fit(s_["player"], "Helvetica-Bold", fs + 1, w - 40))
                c.setFont("Helvetica-Bold", fs); c.setFillGray(0.45); c.drawRightString(x0 + w, yt, f"#{s_['row']}")
                c.setFillGray(0.15); c.setFont("Helvetica", fs)
                for k, ln in enumerate((s_["team"], " · ".join(x for x in (s_["date"], s_["play"]) if x),
                                        (f"{s_['rarity']} · edition size {s_['editions']:,}" if isinstance(s_["editions"], int) else f"{s_['rarity']} · edition size {s_['editions']}") if s_["editions"] else s_["rarity"])):
                    c.drawString(x0, yt - (k + 1) * lh, fit(ln, "Helvetica", fs, w))
            done += len(rs[i0:i0 + per]); end()
        p(f"  {setname}: {len(rs)} cards ({done}/{len(rows)})")
    c.save()
    p(f"Album: {out}  ({total_pages} pages, {os.path.getsize(out)/1e6:.1f} MB, layout {a.layout}, {a.page}{', bleed+crop marks' if BL else ''})")



def main():
    global THR
    ap = argparse.ArgumentParser(description="breakaway-data: NHL Breakaway highlight archive")
    ap.add_argument("command", choices=["all", "fetch", "covers", "build", "package", "videos", "verify", "add", "view", "album"])
    ap.add_argument("refs", nargs="*", help="add: NHL Breakaway URLs or seriesId/tokenId")
    ap.add_argument("--gateway", action="append", default=[], help="IPFS gateway to try first (repeatable), e.g. https://ipfs.io/ipfs/")
    ap.add_argument("--ipfs-only", action="store_true", help="covers/videos: skip NHL Breakaway CDN, use IPFS gateways only")
    ap.add_argument("--concurrency", type=int, default=6, help="parallel requests to NHL Breakaway / CDN (1-8, default 6)")
    ap.add_argument("--refresh", action="store_true", help="fetch: re-download metadata (all, or just --ids) and set listings")
    ap.add_argument("--discover", action="store_true", help="fetch: also add moments from NHL Breakaway set listings that aren't in the card list")
    ap.add_argument("--out-dir", default=os.path.join(OUT, "videos"), help="videos: where to save")
    ap.add_argument("--limit", type=int, help="videos: at most N; fetch --discover: add at most N")
    ap.add_argument("--series", help='videos/album: only this series, e.g. "Series 3"')
    ap.add_argument("--ids", help="fetch/videos: comma-separated card IDs (the id in data/highlights.json)")
    ap.add_argument("--yes", action="store_true", help="videos: don't ask before downloading")
    ap.add_argument("--layout", choices=["1up", "2x2", "3x3"], default="2x2", help="album: cards per page")
    ap.add_argument("--page", choices=["letter", "a4"], default="letter", help="album: paper size")
    ap.add_argument("--set", help="album: only sets whose name contains this")
    ap.add_argument("--player", help="album: only this player (name contains)")
    ap.add_argument("--team", help="album: only this team (name contains)")
    ap.add_argument("--bleed", action="store_true", help="album: add 3 mm bleed and crop marks for print shops")
    ap.add_argument("--out", help="album: output PDF path (default output/album.pdf)")
    ap.add_argument("--max-px", type=int, default=0, help="album: downscale covers to this many pixels (smaller PDF)")
    ap.add_argument("--jpeg", type=int, default=0, help="album: re-encode covers as JPEG at this quality (e.g. 80)")
    a = ap.parse_args(); THR = Throttle(max(1, min(a.concurrency, 8)))
    global GATEWAYS, IPFS_ONLY
    GATEWAYS = [g if g.rstrip("/").endswith("/ipfs") else g.rstrip("/") + "/ipfs/" for g in a.gateway] + DEFAULT_GATEWAYS
    IPFS_ONLY = a.ipfs_only
    if a.command in ("fetch", "all"): cmd_fetch(a)
    if a.command in ("covers", "all"): cmd_covers(a)
    if a.command in ("build", "all"): cmd_build(a)
    if a.command in ("package", "all"): cmd_package(a)
    if a.command == "videos": cmd_videos(a)
    if a.command == "verify": cmd_verify(a)
    if a.command == "add": cmd_add(a)
    if a.command == "view": cmd_view(a)
    if a.command == "album": cmd_album(a)


if __name__ == "__main__":
    main()
